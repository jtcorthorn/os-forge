import csv
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile


SCRIPT = Path(__file__).with_name("extract_sources.py")


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="osforge_extract_test_")
        self.root = Path(self.tmp.name).resolve()
        (self.root / "originals").mkdir()
        self.out = self.root / "_extracted"

    def tearDown(self):
        temp_root = Path(tempfile.gettempdir()).resolve()
        self.assertEqual(self.root.parent, temp_root)
        self.assertTrue(self.root.name.startswith("osforge_extract_test_"))
        self.tmp.cleanup()

    def run_engine(self, rows, extra=()):
        inventory = self.root / "inventory.csv"
        with inventory.open("w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["relative_path", "sha256", "size_bytes"])
            writer.writeheader(); writer.writerows(rows)
        p = subprocess.run([sys.executable, str(SCRIPT), "--root", str(self.root),
                            "--inventory", str(inventory), "--output", str(self.out), *extra],
                           capture_output=True, text=True, timeout=30)
        self.assertEqual(p.returncode, 0, p.stderr)
        return json.loads((self.out / "manifest.json").read_text(encoding="utf-8")), p.stdout

    def row(self, name):
        p = self.root / "originals" / name
        return {"relative_path": "originals/" + name, "sha256": sha(p), "size_bytes": p.stat().st_size}

    def test_text_alias_resume_mismatch_and_safe_html(self):
        a = self.root / "originals" / "a.txt"
        b = self.root / "originals" / "b.txt"
        a.write_text("private alpha\nprivate beta\n", encoding="utf-8")
        b.write_bytes(a.read_bytes())
        h = self.root / "originals" / "c.html"
        h.write_text("<p>visible</p><script>hidden secret</script>", encoding="utf-8")
        bad = self.row("a.txt").copy(); bad["sha256"] = "0" * 64; bad["relative_path"] = "originals/a.txt"
        rows = [self.row("a.txt"), self.row("b.txt"), self.row("c.html")]
        manifest, stdout = self.run_engine(rows)
        self.assertNotIn("private alpha", stdout)
        s = manifest["sources"]
        self.assertEqual([x["status"] for x in s], ["complete"] * 3)
        self.assertEqual(s[1]["alias_of"], s[0]["relative_path"])
        self.assertEqual(s[0]["content"], s[1]["content"])
        text = (self.out / s[2]["content"]).read_text(encoding="utf-8")
        self.assertIn("visible", text)
        self.assertNotIn("hidden secret", text)
        self.assertIn('"locator_kind": "line"', text)
        manifest2, _ = self.run_engine(rows)
        self.assertTrue(manifest2["sources"][0]["resumed"])
        manifest3, _ = self.run_engine([bad])
        self.assertEqual(manifest3["sources"][0]["status"], "hash_mismatch")
        self.assertEqual(a.read_text(encoding="utf-8"), "private alpha\nprivate beta\n")

    def test_docx_pdf_xlsx_and_budget(self):
        docx = self.root / "originals" / "note.docx"
        with zipfile.ZipFile(docx, "w") as z:
            z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Document fact</w:t></w:r></w:p></w:body></w:document>')
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = "Data"; ws["A1"] = "Label"; ws["B1"] = "=1+2"
        xlsx = self.root / "originals" / "book.xlsx"; wb.save(xlsx)
        import fitz
        pdf = self.root / "originals" / "blank.pdf"
        p = fitz.open(); p.new_page(); p.save(pdf); p.close()
        rows = [self.row("note.docx"), self.row("book.xlsx"), self.row("blank.pdf")]
        manifest, _ = self.run_engine(rows)
        s = manifest["sources"]
        self.assertEqual([x["status"] for x in s], ["complete", "complete", "ocr_needed"])
        self.assertEqual(s[2]["metrics"]["pages_without_embedded_text"], 1)
        doc = (self.out / s[0]["content"]).read_text(encoding="utf-8")
        self.assertIn("word/document.xml#paragraph:1", doc)
        book = (self.out / s[1]["content"]).read_text(encoding="utf-8")
        self.assertIn('"formula": "=1+2"', book)
        self.assertIn('"cached_value_missing": true', book)
        manifest2, _ = self.run_engine(rows, ["--max-chars", "3"])
        self.assertEqual(manifest2["sources"][0]["status"], "truncated")
        self.assertNotIn("resumed", manifest2["sources"][0])

    def test_path_escape_rejected(self):
        p = self.root / "originals" / "a.txt"; p.write_text("x")
        row = {"relative_path": "../outside.txt", "sha256": "0" * 64, "size_bytes": 1}
        manifest, _ = self.run_engine([row])
        self.assertEqual(manifest["sources"][0]["status"], "invalid_path")


if __name__ == "__main__":
    unittest.main()
