#!/usr/bin/env python3
"""Bounded, provenance-preserving extraction of a recipient's local source inventory.

Only output under --root is written. Originals are opened read-only.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
from html.parser import HTMLParser
import json
import multiprocessing as mp
import os
from pathlib import Path
import re
import sys
import time
import traceback
import zipfile
from xml.etree import ElementTree as ET

VERSION = "1.0"
TEXT_TYPES = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".log"}
OFFICE_TYPES = {".docx", ".pptx", ".pptm", ".xlsx", ".xlsm"}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".tif", ".tiff", ".webp", ".bmp", ".heic"}
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


class SafeHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.skip = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "svg"):
            self.skip += 1
        elif not self.skip and tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "svg") and self.skip:
            self.skip -= 1
        elif not self.skip and tag in ("p", "div", "li", "tr", "h1", "h2", "h3"):
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def records_for(path: Path, cfg: dict):
    ext = path.suffix.lower()
    budget = cfg["max_chars"]
    max_pages = cfg["max_pages"]
    max_rows = cfg["max_rows"]
    warnings = []
    metrics = {}
    count = 0
    emitted = 0
    truncated = False

    def emit(locator, value, kind="text", extra=None):
        nonlocal emitted, count, truncated
        if value is None:
            return None
        value = str(value)
        if not value.strip():
            return None
        count += 1
        remain = budget - emitted
        if remain <= 0:
            truncated = True
            return None
        if len(value) > remain:
            value = value[:remain]
            truncated = True
        emitted += len(value)
        locator_kind = locator.split(":", 1)[0]
        if "/cell:" in locator: locator_kind = "sheet_cell"
        elif "/paragraph:" in locator: locator_kind = "slide_paragraph" if locator.startswith("slide:") else "document_paragraph"
        rec = {"locator": locator, "locator_kind": locator_kind, "kind": kind, "text": value}
        if extra:
            rec.update(extra)
        return rec

    if ext in TEXT_TYPES or ext in (".html", ".htm"):
        raw = path.read_bytes()
        decoded = raw.decode("utf-8-sig", errors="replace")
        if "\ufffd" in decoded:
            warnings.append("Invalid UTF-8 bytes replaced; verify encoding")
        if ext in (".html", ".htm"):
            parser = SafeHTML()
            parser.feed(decoded)
            decoded = html.unescape("".join(parser.parts))
        metrics["lines"] = decoded.count("\n") + 1
        for n, line in enumerate(decoded.splitlines(), 1):
            item = emit(f"line:{n}", line)
            if item:
                yield item, warnings, metrics
            if truncated:
                break
    elif ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(path)
            if doc.needs_pass:
                raise ProtectedPDF("PDF password required")
            metrics["pages"] = len(doc)
            metrics["pages_without_embedded_text"] = 0
            for i, page in enumerate(doc):
                if i >= max_pages:
                    truncated = True
                    break
                value = page.get_text("text")
                if not value.strip():
                    metrics["pages_without_embedded_text"] += 1
                item = emit(f"page:{i+1}", value)
                if item:
                    yield item, warnings, metrics
                if truncated:
                    break
            doc.close()
        except ImportError:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise ProtectedPDF("PDF password required")
            metrics["pages"] = len(reader.pages)
            metrics["pages_without_embedded_text"] = 0
            for i, page in enumerate(reader.pages):
                if i >= max_pages:
                    truncated = True
                    break
                value = page.extract_text() or ""
                if not value.strip():
                    metrics["pages_without_embedded_text"] += 1
                item = emit(f"page:{i+1}", value)
                if item:
                    yield item, warnings, metrics
                if truncated:
                    break
    elif ext == ".docx":
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist() if n == "word/document.xml" or re.fullmatch(r"word/(headers|footers|footnotes|endnotes|comments)[^/]*\.xml", n)]
            for name in sorted(names):
                root = ET.fromstring(z.read(name))
                for i, p in enumerate(root.findall(".//w:p", NS), 1):
                    value = "".join(n.text or "" for n in p.findall(".//w:t", NS))
                    item = emit(f"{name}#paragraph:{i}", value)
                    if item:
                        yield item, warnings, metrics
                    if truncated:
                        break
                if truncated:
                    break
    elif ext in (".pptx", ".pptm"):
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
            for name in sorted(names, key=lambda n: int(re.search(r"slide(\d+)", n).group(1))):
                slide = int(re.search(r"slide(\d+)", name).group(1))
                if slide > max_pages:
                    truncated = True
                    break
                root = ET.fromstring(z.read(name))
                for i, p in enumerate(root.findall(".//a:p", NS), 1):
                    value = "".join(n.text or "" for n in p.findall(".//a:t", NS))
                    item = emit(f"slide:{slide}/paragraph:{i}", value)
                    if item:
                        yield item, warnings, metrics
                    if truncated:
                        break
                if truncated:
                    break
            metrics["slides"] = len(names)
            if ext == ".pptm":
                warnings.append("Macros not executed")
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        metrics["sheets"] = []
        for ws in formulas.worksheets:
            vws = values[ws.title]
            metrics["sheets"].append({"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column})
            if ws.max_row > max_rows:
                truncated = True
            for row_f, row_v in zip(ws.iter_rows(max_row=min(ws.max_row, max_rows)),
                                    vws.iter_rows(max_row=min(ws.max_row, max_rows))):
                for f, v in zip(row_f, row_v):
                    if f.value is None and v.value is None:
                        continue
                    value = v.value
                    extra = {}
                    if f.data_type == "f":
                        extra["formula"] = str(f.value)
                        if value is None:
                            extra["cached_value_missing"] = True
                    if v.data_type == "e" or f.data_type == "e":
                        extra["formula_error"] = True
                    if value is None and f.value is not None:
                        value = ""
                    if isinstance(value, (str, int, float, bool)):
                        value = str(value)
                    elif value is not None:
                        value = str(value)
                    if not value and "formula" in extra:
                        value = "[cached value missing]"
                    item = emit(f"sheet:{ws.title}/cell:{f.coordinate}", value, "cell", extra)
                    if item:
                        yield item, warnings, metrics
                    if truncated and emitted >= budget:
                        break
                if truncated and emitted >= budget:
                    break
            if truncated and emitted >= budget:
                break
        values.close(); formulas.close()
        warnings.append("Formula cached values only; formulas were not recalculated")
        if ext == ".xlsm":
            warnings.append("Macros not executed")
    else:
        raise Unsupported(ext)

    metrics["text_characters"] = emitted
    metrics["text_records"] = count
    if truncated:
        warnings.append("Truncated by extraction budget; coverage incomplete")
    if ext == ".pdf" and metrics.get("pages_without_embedded_text", 0):
        warnings.append("Pages without embedded text may need OCR")
    yield None, warnings, metrics


class Unsupported(Exception): pass
class ProtectedPDF(Exception): pass


def worker(source: str, out: str, cfg: dict, queue):
    try:
        path, output = Path(source), Path(out)
        status = "complete"
        warnings, metrics = [], {}
        with output.open("w", encoding="utf-8", newline="\n") as f:
            for rec, warnings, metrics in records_for(path, cfg):
                if rec is not None:
                    f.write(json.dumps(rec, ensure_ascii=False, sort_keys=True, default=str) + "\n")
        if "Truncated by extraction budget; coverage incomplete" in warnings:
            status = "truncated"
        elif metrics.get("text_characters", 0) == 0:
            status = "ocr_needed" if path.suffix.lower() == ".pdf" else "no_text"
        queue.put({"status": status, "warnings": warnings, "metrics": metrics})
    except ProtectedPDF as e:
        queue.put({"status": "password_required", "warnings": [str(e)], "metrics": {}})
    except Unsupported as e:
        ext = str(e)
        status = "ocr_needed" if ext in IMAGE_TYPES else "unsupported"
        queue.put({"status": status, "warnings": ["No text parser for format " + ext], "metrics": {}})
    except Exception as e:
        queue.put({"status": "error", "warnings": [type(e).__name__ + ": " + str(e)[:300]], "metrics": {}})


def run(args):
    root = Path(args.root).resolve()
    inventory = Path(args.inventory).resolve()
    output = Path(args.output).resolve()
    if not root.is_dir() or not inventory.is_file():
        raise SystemExit("Root or inventory does not exist")
    if not within(output, root) or output == root:
        raise SystemExit("Output must be a subdirectory inside recipient root")
    if not within(inventory, root):
        raise SystemExit("Inventory must be inside recipient root")
    output.mkdir(parents=True, exist_ok=True)
    content = output / "content"
    content.mkdir(exist_ok=True)
    config = {"max_chars": args.max_chars, "max_pages": args.max_pages, "max_rows": args.max_rows, "timeout": args.timeout}
    sig_payload = {"version": VERSION, "source_sha256": digest(Path(__file__)), "config": config}
    signature = hashlib.sha256(json.dumps(sig_payload, sort_keys=True).encode()).hexdigest()
    prior_path = output / "manifest.json"
    prior = json.loads(prior_path.read_text(encoding="utf-8")) if prior_path.exists() else {}
    prior_entries = {e["relative_path"]: e for e in prior.get("sources", [])} if prior.get("engine_signature") == signature else {}
    with inventory.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    if not rows or not {"relative_path", "sha256"}.issubset(rows[0]):
        raise SystemExit("Inventory needs relative_path,sha256,size_bytes (or bytes)")
    if "size_bytes" not in rows[0] and "bytes" not in rows[0]:
        raise SystemExit("Inventory needs size_bytes or bytes")
    entries = []
    completed = {}
    prefix = Path(args.prefix) if args.prefix else Path()
    if prefix.is_absolute() or ".." in prefix.parts:
        raise SystemExit("Prefix must be relative and remain inside root")
    for row in rows:
        rel = row["relative_path"].replace("\\", "/")
        p = Path(rel)
        if p.is_absolute() or ".." in p.parts or not rel:
            entries.append({"relative_path": rel, "status": "invalid_path"}); continue
        source = (root / prefix / p).resolve()
        if not within(source, root) or within(source, output):
            entries.append({"relative_path": rel, "status": "invalid_path"}); continue
        expected = row["sha256"].lower().strip()
        entry = {"relative_path": rel, "resolved_relative_path": source.relative_to(root).as_posix(),
                 "sha256": expected, "size_bytes": int(row.get("size_bytes") or row.get("bytes") or 0),
                 "format": source.suffix.lower().lstrip(".")}
        if not source.is_file():
            entry.update(status="missing", warnings=["Source file missing"]); entries.append(entry); continue
        actual = digest(source)
        if actual != expected:
            entry.update(status="hash_mismatch", actual_sha256=actual, warnings=["Inventory hash differs from original"]); entries.append(entry); continue
        if source.stat().st_size != entry["size_bytes"]:
            entry.update(status="size_mismatch", actual_size_bytes=source.stat().st_size, warnings=["Inventory size differs from original"]); entries.append(entry); continue
        content_key = expected + "." + entry["format"]
        if content_key in completed:
            base = completed[content_key]
            entry.update({k: base[k] for k in ("status", "warnings", "metrics") if k in base})
            entry["alias_of"] = base["relative_path"]
            if "content" in base: entry["content"] = base["content"]
            entries.append(entry); continue
        old = prior_entries.get(rel)
        artifact = content / (content_key + ".jsonl")
        if old and old.get("sha256") == expected and old.get("status") in ("complete", "truncated", "no_text", "ocr_needed", "unsupported", "password_required") and (old.get("content") is None or artifact.exists()):
            entry.update({k: old[k] for k in ("status", "warnings", "metrics") if k in old})
            if old.get("content"): entry["content"] = old["content"]
            entry["resumed"] = True
        else:
            temp = content / (content_key + ".tmp")
            q = mp.Queue()
            proc = mp.Process(target=worker, args=(str(source), str(temp), config, q))
            proc.start(); proc.join(args.timeout)
            if proc.is_alive():
                proc.terminate(); proc.join()
                result = {"status": "timeout", "warnings": [f"Parser exceeded {args.timeout} seconds"], "metrics": {}}
            elif not q.empty():
                result = q.get()
            else:
                result = {"status": "error", "warnings": [f"Parser exited {proc.exitcode} without result"], "metrics": {}}
            q.close()
            entry.update(result)
            if temp.exists() and result["status"] in ("complete", "truncated", "no_text", "ocr_needed"):
                os.replace(temp, artifact)
                entry["content"] = "content/" + artifact.name
            elif temp.exists():
                temp.unlink()
        completed[content_key] = entry
        entries.append(entry)
    manifest = {"schema_version": 1, "engine_signature": signature, "config": config,
                "root": str(root), "inventory": str(inventory), "prefix": args.prefix or "",
                "sources": entries}
    temp_manifest = output / "manifest.tmp"
    temp_manifest.write_text(json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    os.replace(temp_manifest, prior_path)
    summary = {}
    for e in entries: summary[e["status"]] = summary.get(e["status"], 0) + 1
    print(json.dumps({"manifest": str(prior_path), "source_count": len(entries), "statuses": summary}, sort_keys=True))


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--root", required=True)
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--prefix", default="")
    ap.add_argument("--max-chars", type=int, default=2_000_000)
    ap.add_argument("--max-pages", type=int, default=5000)
    ap.add_argument("--max-rows", type=int, default=10_000)
    ap.add_argument("--timeout", type=int, default=120)
    args = ap.parse_args()
    if min(args.max_chars, args.max_pages, args.max_rows, args.timeout) <= 0:
        ap.error("Budgets and timeout must be positive")
    run(args)


if __name__ == "__main__":
    mp.freeze_support()
    main()
